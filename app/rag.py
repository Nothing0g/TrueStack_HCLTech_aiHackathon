from __future__ import annotations

import hashlib
import re
import shutil
import subprocess
import tempfile
import threading
from pathlib import Path
from typing import Any

from app.config import CHROMA_PATH, EMBEDDING_MODEL, TOP_K, OCR_MAX_PAGES, OCR_DPI, env_bool

_lock = threading.Lock()
_client = None
_collection = None
_embedder = None
_embed_error = None


def _deps():
    """Initialise the persistent Chroma store and local embedding model once."""
    global _client, _collection, _embedder, _embed_error
    with _lock:
        if _collection is not None and _embedder is not None:
            return _collection, _embedder
        try:
            import chromadb
            from sentence_transformers import SentenceTransformer
            CHROMA_PATH.mkdir(parents=True, exist_ok=True)
            _client = chromadb.PersistentClient(path=str(CHROMA_PATH))
            _collection = _client.get_or_create_collection(
                "university_policy_chunks", metadata={"hnsw:space": "cosine"}
            )
            _embedder = SentenceTransformer(EMBEDDING_MODEL)
            _embed_error = None
            return _collection, _embedder
        except Exception as exc:
            _embed_error = f"{type(exc).__name__}: {exc}"
            raise RuntimeError(f"RAG dependencies unavailable: {_embed_error}") from exc


def rag_status() -> dict[str, Any]:
    try:
        collection, _ = _deps()
        return {"status": "ready", "collection_count": collection.count(), "embedding_model": EMBEDDING_MODEL}
    except Exception as exc:
        return {"status": "unavailable", "error": str(exc), "embedding_model": EMBEDDING_MODEL}


def chunk_text(text: str, max_chars: int = 1000, overlap: int = 140) -> list[str]:
    """Split paragraphs into bounded chunks while repeating the active heading."""
    text = re.sub(r"\r", "", text).strip()
    if not text:
        return []
    if max_chars < 100:
        raise ValueError("max_chars must be at least 100")
    overlap = max(0, min(overlap, max_chars // 3))
    blocks = [b.strip() for b in re.split(r"\n\s*\n", text) if b.strip()]
    chunks: list[str] = []
    heading = ""
    for block in blocks:
        lines = block.splitlines()
        first = lines[0].strip() if lines else ""
        if re.match(r"^#{1,4}\s+", first):
            heading = first
            block = "\n".join(lines[1:]).strip()
            if not block:
                continue
        prefix = f"{heading}\n" if heading else ""
        available = max_chars - len(prefix)
        if available <= 0:
            heading = heading[: max_chars - 2]
            prefix = f"{heading}\n" if heading else ""
            available = max_chars - len(prefix)
        if len(block) <= available:
            chunks.append((prefix + block).strip())
            continue
        step = max(1, available - overlap)
        start = 0
        while start < len(block):
            piece = block[start:start + available].strip()
            if piece:
                chunks.append((prefix + piece).strip())
            if start + available >= len(block):
                break
            start += step
    return chunks

def _ocr_pdf_pages(path: Path, page_indices: list[int]) -> dict[int, str]:
    """OCR selected zero-based PDF page indices with local Poppler + Tesseract.

    OCR is deliberately bounded: for long scanned documents, use an OCR-enabled copy
    or upload relevant text-readable excerpts rather than freezing the API for minutes.
    Returns a mapping keyed by one-based physical PDF page number.
    """
    if not page_indices:
        return {}
    if len(page_indices) > OCR_MAX_PAGES:
        raise ValueError(
            f"Scanned PDF needs OCR on {len(page_indices)} pages, above OCR_MAX_PAGES={OCR_MAX_PAGES}. "
            "OCR it locally first or upload a text-readable PDF."
        )
    if not shutil.which("pdftoppm") or not shutil.which("tesseract"):
        raise ValueError("Scanned PDF requires Poppler (pdftoppm) and Tesseract. Install them or upload a text-readable PDF.")
    output: dict[int, str] = {}
    with tempfile.TemporaryDirectory(prefix="truestack-ocr-") as tmp:
        base = Path(tmp)
        for zero_index in page_indices:
            number = zero_index + 1
            prefix = base / f"page-{number}"
            try:
                subprocess.run(
                    ["pdftoppm", "-f", str(number), "-l", str(number), "-singlefile", "-r", str(OCR_DPI), "-png", str(path), str(prefix)],
                    check=True, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE, timeout=30,
                )
                image_path = prefix.with_suffix(".png")
                result = subprocess.run(
                    ["tesseract", str(image_path), "stdout", "-l", "eng", "--psm", "6"],
                    check=True, capture_output=True, text=True, timeout=35,
                )
                text = result.stdout.strip()
                if text:
                    output[number] = text
            except subprocess.TimeoutExpired as exc:
                raise ValueError(f"OCR timed out on PDF page {number}.") from exc
            except (subprocess.CalledProcessError, OSError) as exc:
                raise ValueError(f"OCR failed on PDF page {number}: {type(exc).__name__}.") from exc
    return output


def read_document(path: Path) -> tuple[str, list[dict[str, Any]]]:
    """Read PDF/DOCX/Markdown/TXT while retaining real PDF page anchors.

    For image-only pages, OCR is used when OCR_ENABLED is true (default true).
    Empty or unusable PDFs are rejected rather than represented as successfully indexed.
    """
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        text = path.read_text(encoding="utf-8")
        return text, []
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        pages: list[dict[str, Any]] = []
        blank_pages: list[int] = []
        for number, page in enumerate(reader.pages, start=1):
            body = (page.extract_text() or "").strip()
            if len(body) < 25:
                blank_pages.append(number)
            pages.append({"page": number, "text": body})
        if blank_pages and env_bool("OCR_ENABLED", True):
            ocr = _ocr_pdf_pages(path, [n - 1 for n in blank_pages])
            for item in pages:
                if item["page"] in ocr and len(ocr[item["page"]]) > len(item["text"]):
                    item["text"] = ocr[item["page"]]
        combined = "\n\n".join(f"[PAGE {item['page']}]\n{item['text']}" for item in pages if item["text"].strip())
        if len(combined.strip()) < 30:
            raise ValueError("No usable text could be extracted. If this is a scanned PDF, enable OCR and install Tesseract.")
        return combined, pages
    if suffix == ".docx":
        from docx import Document
        document = Document(str(path))
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        table_text = []
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    table_text.append(" | ".join(cells))
        text = "\n\n".join(paragraphs + table_text)
        if len(text.strip()) < 30:
            raise ValueError("No usable text was found in the DOCX document.")
        return text, []
    raise ValueError("Unsupported file type. Allowed: .pdf, .docx, .md, .txt")


def remove_document(doc_id: str) -> None:
    """Remove all vectors for a document; used to recover from a failed ingestion transaction."""
    collection, _ = _deps()
    existing = collection.get(where={"doc_id": str(doc_id)}, include=[])
    ids = existing.get("ids") or []
    if ids:
        collection.delete(ids=ids)


def ingest_path(text: str, metadata: dict[str, Any], source_path: str = "", reset: bool = False) -> int:
    collection, embedder = _deps()
    page_sections = re.split(r"(?=\[PAGE \d+\])", text) if re.search(r"\[PAGE \d+\]", text) else [text]
    chunks: list[str] = []
    for page_text in page_sections:
        page_text = page_text.strip()
        if not page_text:
            continue
        anchor = re.match(r"\[PAGE (\d+)\]", page_text)
        anchor_text = anchor.group(0) if anchor else ""
        body = page_text[anchor.end():].strip() if anchor else page_text
        for chunk in chunk_text(body):
            if chunk.strip():
                chunks.append((anchor_text + "\n" if anchor_text else "") + chunk.strip())
    if not chunks:
        raise ValueError("No extractable text found; document was not indexed.")

    doc_id = str(metadata["doc_id"])
    ids = [hashlib.sha256(f"{doc_id}:{i}:{chunk}".encode("utf-8")).hexdigest() for i, chunk in enumerate(chunks)]
    metadatas = []
    for i, chunk in enumerate(chunks):
        item = {
            key: str(metadata.get(key, ""))
            for key in ["doc_id", "title", "issuer", "doc_type", "version", "effective_from", "effective_to",
                        "supersedes", "scope_programmes", "scope_batches", "provenance", "synthetic"]
        }
        item["authority_level"] = int(metadata.get("authority_level", 5))
        item["chunk_index"] = i
        page_match = re.search(r"^\[PAGE (\d+)\]", chunk)
        item["page"] = int(page_match.group(1)) if page_match else 0
        without_page = re.sub(r"^\[PAGE \d+\]\s*", "", chunk.strip(), flags=re.I)
        section_match = re.search(
            r"^(?:#{1,4}\s+|(?:Section|Clause)\s+|(?:R|Regulation)\.?\s*\d+\s*\([A-Za-z0-9]+\)\.?\d+|\d+(?:\.\d+){1,3})[^\n]{0,140}",
            without_page, re.I | re.M,
        )
        item["section"] = section_match.group(0).strip().lstrip("# ") if section_match else ""
        metadatas.append(item)
    embeddings = embedder.encode(chunks, normalize_embeddings=True).tolist()
    if reset:
        remove_document(doc_id)
    collection.upsert(ids=ids, documents=chunks, embeddings=embeddings, metadatas=metadatas)
    return len(chunks)


def search(query: str, top_k: int | None = None) -> list[dict[str, Any]]:
    collection, embedder = _deps()
    q = embedder.encode([query], normalize_embeddings=True).tolist()
    count = collection.count()
    if not count:
        return []
    results = collection.query(query_embeddings=q, n_results=min(top_k or TOP_K, count), include=["documents", "metadatas", "distances"])
    output = []
    for index, text in enumerate((results.get("documents") or [[]])[0]):
        metadata = (results.get("metadatas") or [[]])[0][index] or {}
        distance = (results.get("distances") or [[]])[0][index]
        output.append({"text": text, "metadata": metadata, "score": float(1-distance) if distance is not None else None})
    return output


def search_document(doc_id: str) -> list[dict[str, Any]]:
    """Return indexed chunks for one exact source ID, preserving stored provenance."""
    collection, _ = _deps()
    result = collection.get(where={"doc_id": str(doc_id)}, include=["documents", "metadatas"])
    documents = result.get("documents") or []
    metadatas = result.get("metadatas") or []
    ids = result.get("ids") or []
    out = []
    for i, body in enumerate(documents):
        metadata = metadatas[i] or {} if i < len(metadatas) else {}
        out.append({"id": ids[i] if i < len(ids) else "", "text": body, "metadata": metadata, "score": None})
    return sorted(out, key=lambda x: int(x.get("metadata", {}).get("chunk_index", 0)))


def delete_document(doc_id: str) -> int:
    """Delete all vectors for a document; used when ingestion fails after indexing."""
    collection, _ = _deps()
    result = collection.get(where={"doc_id": str(doc_id)}, include=["metadatas"])
    ids = result.get("ids") or []
    if ids:
        collection.delete(ids=ids)
    return len(ids)
